#!/usr/bin/env python3
"""Build an armored drill book (PDF + DOCX) from a BOOK spec. Usage: python3 build_book.py book_itsm"""
import sys, os, random, importlib, re, datetime
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Flowable
from reportlab.lib.enums import TA_CENTER
from xml.sax.saxutils import escape as e
from docx import Document
from docx.shared import Pt

HERE=os.path.dirname(os.path.abspath(__file__)); OUT=os.path.join(os.path.dirname(HERE),"editions"); os.makedirs(OUT,exist_ok=True)
mod=importlib.import_module(sys.argv[1]); B=mod.BOOK
MOLD={"DROP":"dropped element","STEM":"wrong stem number","ADJ":"adjacent category","INV":"inverted operation","ABS":"absolute language","CONF":"confusion set","CAMO":"answers the camouflage"}
L="ABCD"; random.seed(hash(B["title"])%10007)
items=[]
slots=list(range(4)); random.shuffle(slots)
for n,it in enumerate(B["items"]):
    keys=[x for x in it["opts"] if x[1]=="KEY"]; dis=[x for x in it["opts"] if x[1]!="KEY"]; random.shuffle(dis)
    pos=slots[n%4]; o=[None]*4
    o[pos]=keys[0]; rest=[i for i in range(4) if i!=pos]
    if len(keys)>1: o[rest.pop(random.randrange(len(rest)))]=keys[1]
    for i in rest: o[i]=dis.pop()
    it2=dict(it); it2["opts"]=o; it2["keys"]=[L[i] for i,x in enumerate(o) if x[1]=="KEY"]; items.append(it2)
dist={k:sum(1 for it in items if k in it["keys"]) for k in L}
today=datetime.date.today().isoformat()
DISCLAIMER=B.get("disclaimer") or (f"{B['platform']} is a trademark of its owner. This book is an independent, unofficial study aid. It is not affiliated with, "
 f"sponsored by, or endorsed by {B['platform']} or any employer of the author. Every item is original and written from the platform's public documentation; "
 "no certification exam question is reproduced or paraphrased, and no customer instance, configuration or data is described.")
if B.get("intro"): intro=B["intro"]

def S(**k):
    d=dict(fontName="Helvetica",fontSize=9.6,leading=12.2); d.update(k); return ParagraphStyle("x",**d)
st=S(); small=S(fontSize=8.6,leading=10.8); h1=S(fontName="Helvetica-Bold",fontSize=22,leading=27); h2=S(fontName="Helvetica-Bold",fontSize=12.5,leading=15,spaceBefore=10,spaceAfter=5)
class Rule(Flowable):
    def __init__(s,w=6.9*inch): super().__init__(); s.w=w; s.height=4
    def draw(s): c=s.canv; c.setLineWidth(0.6); c.setDash(2,2); c.line(0,2,s.w,2)
def box(fl):
    t=Table([[fl]],colWidths=[6.95*inch]); t.setStyle(TableStyle([("BOX",(0,0),(-1,-1),0.8,colors.black),("LEFTPADDING",(0,0),(-1,-1),7),("RIGHTPADDING",(0,0),(-1,-1),7),("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)])); return t
bl=lambda w:"_"*w
def q(n,it):
    fl=[Paragraph(f"<b>{n}.</b> <b>{e(it['concept'])}</b>",st),Spacer(1,2),Paragraph(e(it["stem"]),st),Spacer(1,3)]
    for i,(t,_) in enumerate(it["opts"]): fl.append(Paragraph(f"&nbsp;&nbsp;<b>{L[i]}.</b> {e(t)}",st))
    fl+=[Spacer(1,4),Rule(),Spacer(1,2),Paragraph(f"Target word: {bl(22)} &nbsp; The one fact that decides: {bl(46)}",small),
         Paragraph(f"Answer: {bl(8)} &nbsp;&nbsp; [&nbsp;&nbsp;&nbsp;] sure &nbsp;&nbsp; [&nbsp;&nbsp;&nbsp;] unsure &nbsp;&nbsp; [&nbsp;&nbsp;&nbsp;] guess",small)]
    return [box(fl),Spacer(1,6)]
def k(n,it):
    fl=[Paragraph(f"<b>{n}.</b> <b>{e(it['concept'])}</b> — <b>{' and '.join(it['keys'])}</b>",st),Spacer(1,2),
        Paragraph(f"<b>Rule.</b> {e(it['rule'])}",st),Paragraph(f"<b>Target word.</b> {e(it['target'])} &nbsp;&nbsp; <b>Decides.</b> {e(it['deciding'])}",st)]
    camo=" · ".join(f"<strike>{e(c)}</strike>" for c in it.get("camo",[])) or "—"
    fl.append(Paragraph(f"<b>Camouflage, struck.</b> {camo}",st))
    rows=[[Paragraph(f"<b>{L[i]}</b>" if m=="KEY" else L[i],small),Paragraph("<b>KEY</b>" if m=="KEY" else MOLD[m],small),Paragraph(e(t),small)] for i,(t,m) in enumerate(it["opts"])]
    tb=Table(rows,colWidths=[0.25*inch,1.15*inch,5.2*inch]); tb.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.3,colors.grey),("VALIGN",(0,0),(-1,-1),"TOP"),("TOPPADDING",(0,0),(-1,-1),1.5),("BOTTOMPADDING",(0,0),(-1,-1),1.5)]))
    fl+=[Spacer(1,3),tb]; return [box(fl),Spacer(1,6)]

SERIES="Joshua Says"; TAG="For people who already know the easy version."; AUTHOR="Joshua"
story=[Spacer(1,1.2*inch),Paragraph(e(SERIES),S(fontName="Helvetica-Bold",fontSize=13,leading=16)),Spacer(1,4),Paragraph(e(B["title"]),h1),Spacer(1,6),Paragraph(e(TAG),S(fontName="Helvetica-Oblique",fontSize=12,leading=15)),Spacer(1,10),Paragraph(e(B["subtitle"]),S(fontSize=11,leading=14)),Spacer(1,14),Paragraph(f"by {AUTHOR}",S(fontSize=11)),Spacer(1,14),
 Paragraph("<b>V1.</b> Written from the platform's public documentation as of the date below. Check every key against the source area named in <i>Check the work</i> before relying on it; corrections are dated and kept public.",st),Spacer(1,6),
 Paragraph((f"{len(items)} items · " if items else "")+f"V1 · {today}",st),Spacer(1,1.8*inch),Paragraph(e(DISCLAIMER),small),PageBreak(),
 Paragraph("How to run it" if B["items"] else "How to use it",h2)]
intro=[["1","<b>Read the target word first.</b> Write it on the line before you read the options. The platform disguises the question in the noun it uses — case or incident, skill or agent, queue or division, total or per unit."],
 ["2","<b>Name the one fact that decides.</b> Everything else in the stem is camouflage: the deadline, the executive, the numbers, the object's name."],
 ["3","<b>Principle before configuration.</b> Say what the platform construct is <i>for</i>, then pick the construct."],
 ["4","<b>Confusion sets are the platform's favourite trap.</b> Every wrong option is a real thing. The test is exact role-matching, not recognition."],
 ["5","<b>Mark confidence honestly.</b> Sure-and-wrong is a wrong mental model, and it is the only thing worth finding."],
 ["6","<b>Grade in Part B.</b> Each distractor carries the error it encodes. Name the mold you fell for; naming it defeats it."]]
t=Table([[Paragraph(a,S(fontName="Helvetica-Bold")),Paragraph(b,st)] for a,b in intro],colWidths=[0.3*inch,6.6*inch]); t.setStyle(TableStyle([("BOX",(0,0),(-1,-1),0.8,colors.black),("VALIGN",(0,0),(-1,-1),"TOP")]))
story+=[t,Spacer(1,10),Paragraph("Check the work — sources",h2),Paragraph("Every item is written from public documentation. Before trusting a key, check it against the source area named here; corrections are dated and kept public.",st),Spacer(1,4)]
srows=[[Paragraph(f"<b>{e(a)}</b>",small),Paragraph(e(b),small)] for a,b in B["sources"]]
ts=Table(srows,colWidths=[2.4*inch,4.5*inch]); ts.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.3,colors.grey),("VALIGN",(0,0),(-1,-1),"TOP")]))
story+=[ts,Spacer(1,8)]+([Paragraph("The seven molds: dropped element · wrong stem number · adjacent category · inverted operation · absolute language · confusion set · camouflage (the option that answers a stem fact that was never the question).",small)] if items else [])+[PageBreak()]
for sec in B.get("sections",[]):
    story.append(Paragraph(e(sec["h"]),h2))
    for para in sec["p"]:
        if isinstance(para,list):
            tb=Table([[Paragraph(e(c),small) for c in row] for row in para],colWidths=[6.9*inch/len(para[0])]*len(para[0])); tb.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.3,colors.grey),("VALIGN",(0,0),(-1,-1),"TOP"),("BACKGROUND",(0,0),(-1,0),colors.whitesmoke)])); story+=[tb,Spacer(1,6)]
        else: story+=[Paragraph(e(para),st),Spacer(1,4)]
if B.get("sections") and items: story.append(PageBreak())
if items:
    story.append(Paragraph("Part A — Items",h2))
    for n,it in enumerate(items,1): story+=q(n,it)
if items:
    story+=[PageBreak(),Spacer(1,3*inch),Paragraph("STOP",S(fontName="Helvetica-Bold",fontSize=28,leading=34,alignment=TA_CENTER)),Spacer(1,8),Paragraph("Answers ahead. Finish Part A and mark every confidence box, then turn.",S(alignment=TA_CENTER,fontSize=11)),PageBreak(),Paragraph("Part B — Key and dissection",h2)]
    for n,it in enumerate(items,1): story+=k(n,it)
tri=[["",Paragraph("<b>Right</b>",st),Paragraph("<b>Wrong</b>",st)],[Paragraph("<b>Sure</b>",st),Paragraph("Owned. Bank it.",small),Paragraph("<b>Wrong model. Top priority.</b> Say the rule aloud; find the construct in the documentation.",small)],
 [Paragraph("<b>Unsure</b>",st),Paragraph("Recognition, not mastery. Redo from the rule.",small),Paragraph("Not learned. Read the rule, then the source area.",small)],[Paragraph("<b>Guess</b>",st),Paragraph("Luck. Treat as not learned.",small),Paragraph("Not learned.",small)]]
tt=Table(tri,colWidths=[0.9*inch,3*inch,3*inch],rowHeights=[0.3*inch,1.2*inch,1.2*inch,1.2*inch]); tt.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.8,colors.black),("VALIGN",(0,0),(-1,-1),"TOP")]))
if items: story+=[PageBreak(),Paragraph("Part C — Triage",h2),Paragraph("Sort every item by the two marks you made. Write item numbers in the cells.",st),Spacer(1,6),tt,Spacer(1,10),
 Paragraph("<b>Mold tally.</b> dropped element ____ · wrong stem number ____ · adjacent category ____ · inverted operation ____ · absolute language ____ · confusion set ____ · camouflage ____ . The largest number is the pattern to drill.",st),
 Spacer(1,16),Paragraph(e(DISCLAIMER),small),Spacer(1,6),Paragraph(f"© {today[:4]} Joshua. From Joshua's mind — that is the signature.",small)]
slug=re.sub(r"[^A-Za-z0-9]+","_",B["title"]).strip("_")
pdf=os.path.join(OUT,f"{slug}.pdf")
def foot(c,d): c.setFont("Helvetica",8); c.drawRightString(letter[0]-0.6*inch,0.4*inch,f"Joshua Says · {B['title']} · V1 · {d.page}")
SimpleDocTemplate(pdf,pagesize=letter,leftMargin=0.6*inch,rightMargin=0.6*inch,topMargin=0.6*inch,bottomMargin=0.6*inch,title=B["title"]).build(story,onFirstPage=foot,onLaterPages=foot)

# DOCX (KDP manuscript)
d=Document(); d.styles["Normal"].font.name="Calibri"; d.styles["Normal"].font.size=Pt(10.5)
d.add_paragraph("Joshua Says"); d.add_heading(B["title"],0); d.add_paragraph("For people who already know the easy version."); d.add_paragraph(B["subtitle"]); d.add_paragraph("by Joshua"); d.add_paragraph(f"V1 · {today}"); d.add_paragraph(DISCLAIMER)
d.add_page_break(); d.add_heading("How to run it" if items else "How to use it",1)
for a,b in intro: d.add_paragraph(re.sub(r"<[^>]+>","",b),style="List Number")
d.add_heading("Check the work — sources",1)
for a,b in B["sources"]: d.add_paragraph(f"{a} — {b}",style="List Bullet")
for sec in B.get("sections",[]):
    d.add_heading(sec["h"],1)
    for para in sec["p"]:
        if isinstance(para,list):
            t=d.add_table(rows=len(para),cols=len(para[0])); t.style="Table Grid"
            for r,row in enumerate(para):
                for c,cell in enumerate(row): t.cell(r,c).text=cell
        else: d.add_paragraph(para)
if items: d.add_page_break(); d.add_heading("Part A — Items",1)
for n,it in enumerate(items,1):
    d.add_paragraph(f"{n}. {it['concept']}").runs[0].bold=True; d.add_paragraph(it["stem"])
    for i,(t,_) in enumerate(it["opts"]): d.add_paragraph(f"{L[i]}. {t}")
    d.add_paragraph("Target word: ____________   The one fact that decides: ______________________________"); d.add_paragraph("Answer: ____   [ ] sure  [ ] unsure  [ ] guess")
if items: d.add_page_break(); d.add_heading("STOP — answers ahead",1); d.add_page_break(); d.add_heading("Part B — Key and dissection",1)
for n,it in enumerate(items,1):
    d.add_paragraph(f"{n}. {it['concept']} — {' and '.join(it['keys'])}").runs[0].bold=True
    d.add_paragraph(f"Rule. {it['rule']}"); d.add_paragraph(f"Target word. {it['target']}   Decides. {it['deciding']}")
    if it.get("camo"): d.add_paragraph("Camouflage, struck: "+" · ".join(it["camo"]))
    for i,(t,m) in enumerate(it["opts"]): d.add_paragraph(f"{L[i]}  {'KEY' if m=='KEY' else MOLD[m]}  —  {t}")
if items: d.add_page_break(); d.add_heading("Part C — Triage",1)
if items: d.add_paragraph("Sure + right = owned · Sure + wrong = wrong model, top priority · Unsure + right = recognition, not mastery · Unsure/guess + wrong = not learned.")
if items: d.add_paragraph("Mold tally: dropped element __ · wrong stem number __ · adjacent category __ · inverted operation __ · absolute language __ · confusion set __ · camouflage __")
d.add_paragraph(DISCLAIMER)
docx=os.path.join(OUT,f"{slug}.docx"); d.save(docx)
# scan for forbidden terms
txt=open(sys.argv[1]+".py",encoding="utf-8").read().lower()
hits=[w for w in ["wgu","western governors","enrollment team","uqa","lindquist","el.3079","eriklindquist"] if w in txt]
print(f"{B['title']}: {len(items)} items · keys {dist} · scan {hits or 'clean'} · {os.path.basename(pdf)}")
